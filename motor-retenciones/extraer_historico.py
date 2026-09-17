# -*- coding: utf-8 -*-
"""Extrae el historico 2026 de los dos Excel a un dataset normalizado (JSON).

No interpreta ni corrige nada: solo lee tal cual esta, marcando que fila es
dato, que fila es encabezado de mes y que fila es subtotal de quincena.
"""
import datetime, json, re, sys
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent.parent
SALIDA = Path(__file__).resolve().parent / "historico.json"


def norm(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v if v not in ("", "-", "--") else None
    return v


def fecha(v):
    """Devuelve ISO o el texto crudo si no es una fecha parseable."""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    return norm(v)


# --- marcas por color, segun la leyenda de la propia planilla (fila 213-215) ---
VERDE = "FF92D050"          # "FC USD - TC del dia"
ROJO_FUENTE = "FFFF0000"    # "ROJO = No corresponde / Anulada"
AMARILLOS = {"FFFFFF00", "FFFFFF99"}   # sin leyenda en la planilla


def _fill(ws, fila, col):
    c = ws.cell(row=fila, column=col)
    if c.fill is None or c.fill.fill_type != "solid":
        return None
    sc = c.fill.start_color
    if sc.type == "rgb":
        return sc.rgb
    if sc.type == "theme":
        return f"theme{sc.theme}"
    return None


def marcas(ws, fila):
    """Marcas de negocio codificadas por color, segun la leyenda de la planilla.

    El verde y el azul pintan la columna del monto (E); el amarillo, la del
    proveedor (I); el rojo es color de fuente sobre toda la fila.
    """
    monto = _fill(ws, fila, 5)
    prov = _fill(ws, fila, 9)
    f = ws.cell(row=fila, column=9).font
    rojo = bool(f and f.color and f.color.type == "rgb" and f.color.rgb == ROJO_FUENTE)
    return {
        "anulada": rojo,                       # "No corresponde / Anulada"
        "fc_usd": monto == VERDE,              # "FC USD - TC del dia"
        "res_1556": monto == "theme8",         # "RETENCION IVA Y SUSS *RES 1556"
        "resaltado": prov in AMARILLOS,        # sin leyenda: a preguntar
    }


def es_encabezado_mes(row):
    """Fila con solo una fecha en la columna A -> arranca un periodo."""
    return isinstance(row[0], datetime.datetime) and all(
        norm(c) is None for c in row[1:6]
    )


def cuit(v):
    if v is None:
        return None
    s = re.sub(r"\D", "", str(v))
    return s.zfill(11) if s else None


def extraer_ganancias(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Ret. Realizadas 2026"]
    filas, subtotales = [], []
    periodo = None
    for i, r in enumerate(ws.iter_rows(min_row=1, max_row=3000, values_only=True), 1):
        r = list(r) + [None] * (14 - len(r))
        if all(norm(c) is None for c in r):
            continue
        if es_encabezado_mes(r):
            periodo = r[0].strftime("%Y-%m")
            continue
        if isinstance(r[0], str) and r[0].strip().lower().startswith("fecha"):
            continue
        if isinstance(r[6], str):                      # "1Q" / "2Q"
            subtotales.append({"fila": i, "periodo": periodo,
                               "quincena": r[6].strip(), "total": r[7]})
            continue
        if not isinstance(r[4], (int, float)):         # sin Monto Neto -> no es dato
            continue
        filas.append({
            "fila": i, "periodo": periodo,
            "fecha_factura": fecha(r[0]), "nro_factura": norm(r[1]), "pzf": norm(r[2]),
            "regimen": norm(r[3]), "monto_neto": r[4], "base_imponible": r[5],
            "alicuota": r[6], "retencion": r[7],
            "razon_social": norm(r[8]), "cuit": cuit(r[9]),
            "fecha_retencion": fecha(r[10]), "nro_certificado": norm(r[11]),
            "ret_iva_966": norm(r[12]), "ret_suss": norm(r[13]),
            **marcas(ws, i),
        })
    wb.close()
    return filas, subtotales


def extraer_caba(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Retenciones"]
    filas = []
    periodo = None
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        r = list(r) + [None] * (13 - len(r))
        if all(norm(c) is None for c in r):
            continue
        if es_encabezado_mes(r):
            periodo = r[0].strftime("%Y-%m")
            continue
        if isinstance(r[0], str) and r[0].strip().lower().startswith("fecha"):
            continue
        if not isinstance(r[5], (int, float)):
            continue
        filas.append({
            "fila": i, "periodo": periodo,
            "fecha_factura": fecha(r[0]), "nro_factura": norm(r[1]), "pzf": norm(r[2]),
            "comp": norm(r[3]), "fc_tipo": norm(r[4]),
            "monto_neto": r[5], "alicuota": r[6], "retencion": r[7],
            "razon_social": norm(r[8]), "cuit": cuit(r[9]),
            "fecha_retencion": fecha(r[10]), "nro_certificado": norm(r[11]),
            "padron_crudo": norm(r[12]),
            **marcas(ws, i),
        })
    wb.close()
    return filas


def extraer_fc_c(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["FC C "]
    filas, periodo = [], None
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        r = list(r) + [None] * (9 - len(r))
        if all(norm(c) is None for c in r):
            continue
        if es_encabezado_mes(r):
            periodo = r[0].strftime("%Y-%m")
            continue
        if isinstance(r[0], str) and r[0].strip().lower().startswith("fecha"):
            continue
        if not isinstance(r[4], (int, float)):
            continue
        filas.append({
            "fila": i, "periodo": periodo, "fecha_factura": fecha(r[0]),
            "nro_factura": norm(r[1]), "pzf": norm(r[2]), "regimen": norm(r[3]),
            "monto_neto": r[4], "razon_social": norm(r[5]), "cuit": cuit(r[6]),
            "fecha_analisis": fecha(r[7]), "motivo": norm(r[8]),
        })
    wb.close()
    return filas


def main():
    gan, subt = extraer_ganancias(BASE / "Retenciones_Ganancias_Practicadas_2026.xlsx")
    caba = extraer_caba(BASE / "Retenciones_CABA_Agente_Recaudacion_2026.xlsx")
    fcc = extraer_fc_c(BASE / "Retenciones_Ganancias_Practicadas_2026.xlsx")

    data = {"ganancias": gan, "subtotales_ganancias": subt, "caba": caba, "fc_c": fcc}
    SALIDA.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")

    print(f"ganancias : {len(gan):3} filas   periodos {sorted({f['periodo'] for f in gan})}")
    print(f"caba      : {len(caba):3} filas   periodos {sorted({f['periodo'] for f in caba})}")
    print(f"fc c      : {len(fcc):3} filas")
    print(f"subtotales: {len(subt)}")
    print(f"-> {SALIDA}")

    sin_periodo = [f["fila"] for f in gan + caba if not f["periodo"]]
    if sin_periodo:
        print("AVISO filas sin periodo:", sin_periodo)
    for nombre, filas in (("ganancias", gan), ("caba", caba)):
        print(f"{nombre:9} marcas:",
              {m: sum(1 for f in filas if f[m]) for m in ("anulada", "fc_usd", "res_1556", "resaltado")})
    sin_cuit = [f["fila"] for f in gan if not f["cuit"]]
    if sin_cuit:
        print("AVISO ganancias sin cuit:", sin_cuit)


if __name__ == "__main__":
    main()
